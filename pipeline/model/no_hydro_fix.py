"""
no_hydro_fix.py
===============
NO-spezifische Hydro-Korrektur für PSOpen_NO, gelesen aus den
PEMMDB-Bietzonen-Dateien der norwegischen Zonen (NOM1, NON1, NOS0):

- Aggregierte Speicherkapazität für PSOpen_NO aus mehreren Bietzonen
- Wöchentliche Erzeugungs-Constraints (min/max) für PSOpen_NO
- Wöchentliche Reservoir-Level-Constraints (min/max State-of-Charge)

Die Excel-Dateien werden unterhalb von config.ACCUMULATED_DIR gesucht
(liegen in "Climate Data/Hydro Inflows/NO constraint approach/").

Aktiv, wenn config.NO_HYDRO_FIX bzw. --no-hydro-fix != "off".
"""

import logging
import warnings
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import pandas as pd
import pypsa
from openpyxl import load_workbook

import config

logger = logging.getLogger(__name__)

warnings.filterwarnings(
    "ignore",
    message=r"Defined names for sheet index .* cannot be located",
    category=UserWarning,
)

NO_BZONE_CODES = ["NOM1", "NON1", "NOS0"]
HYDRO_ZONE_CONFIGS = [
    {"zone": "NO", "codes": NO_BZONE_CODES, "storage_name": "PSOpen_NO"},
]


def _find_no_hydro_files(search_dir: Path, codes: Optional[List[str]] = None) -> List[Path]:
    """Findet relevante norwegische PEMMDB-Hydro-Dateien unterhalb von search_dir."""
    codes = codes or NO_BZONE_CODES
    candidates: List[Path] = []
    seen = set()

    patterns = [
        "**/PEMMDB_{code}_Hydro_Inflow_2030.xlsx",
        "**/PEMMDB_{code}_Hydro_Inflow_2030.csv",
        "**/PEMMDB_{code}_Hydro Inflow_2030.xlsx",
        "**/PEMMDB_{code}_Hydro Inflow_2030.csv",
        "**/*{code}*Hydro*Inflow*.xlsx",
        "**/*{code}*Hydro*Inflow*.csv",
        "**/*{code}*Hydro*Inflow*.*",
        "**/*{code}*accumulated*.csv",
    ]

    for code in codes:
        for pattern in patterns:
            for path in search_dir.rglob(pattern.format(code=code)):
                if path.is_file() and path not in seen:
                    candidates.append(path)
                    seen.add(path)

    return candidates


def _coerce_numeric(value) -> Optional[float]:
    """Konvertiert einen Zellwert in einen numerischen Float, falls möglich."""
    if value is None or pd.isna(value):
        return None
    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).strip()
    if not text:
        return None
    text = text.replace(" ", "")
    if text.startswith("(") and text.endswith(")"):
        text = text[1:-1]
    try:
        return float(text.replace(".", "").replace(",", "."))
    except ValueError:
        return None


def _read_hydro_input_tables(path: Path) -> List[pd.DataFrame]:
    """Liest die Hydro-Input-Datei robust als Excel-Workbook oder CSV ein.

    Sortiert Sheets so, dass 'Pump storage - Open Loop' zuerst verarbeitet wird,
    da andere Sheets (Pondage, Reservoir) ebenfalls reservoir-Marker haben, aber
    Kapazitätswerte von 0 für NO enthalten können.
    """
    if path.suffix.lower() == ".csv":
        return [pd.read_csv(path, header=None, low_memory=False)]

    workbook = load_workbook(path, read_only=True, data_only=True)

    def _sheet_priority(ws) -> int:
        name = ws.title.lower()
        if "pump" in name and "open" in name:
            return 0
        if "pump" in name:
            return 1
        return 2

    sheets = sorted(workbook.worksheets, key=_sheet_priority)
    return [pd.DataFrame(list(sheet.iter_rows(values_only=True))) for sheet in sheets]


def _find_label_row(df: pd.DataFrame, labels: List[str]) -> Tuple[Optional[int], Dict[str, int]]:
    """Sucht eine Zeile, die alle gesuchten Labels enthält."""
    for row_idx, row in df.iterrows():
        row_vals = [str(val).strip().lower() for val in row.tolist()]
        matches = {}
        for label in labels:
            label_lower = label.lower()
            for col_idx, val in enumerate(row_vals):
                if label_lower in val:
                    matches[label] = col_idx
                    break
        if all(label in matches for label in labels):
            return row_idx, matches
    return None, {}


def _find_numeric_value(df: pd.DataFrame, labels: List[str]) -> Optional[float]:
    """Sucht in der Tabelle nach einem numerischen Wert, dessen Zeile die Labels enthält."""
    for row_idx, row in df.iterrows():
        row_vals = [str(val).strip().lower() for val in row.tolist() if pd.notna(val)]
        if not any(label.lower() in val for label in labels for val in row_vals):
            continue

        for col_idx, val in enumerate(row.tolist()):
            numeric = _coerce_numeric(val)
            if numeric is not None:
                return numeric
    return None


def load_no_hydro_weekly_constraints(search_dir: Path, codes: Optional[List[str]] = None) -> Tuple[pd.DataFrame, Dict[str, float]]:
    """
    Liest weekly Generation-Bounds aus den norwegischen PEMMDB-Hydro-Dateien,
    summiert sie über alle Bietzonen und gibt sie als DataFrame zurück.
    """
    files = _find_no_hydro_files(search_dir, codes=codes)

    if not files:
        logger.warning("Keine norwegischen Hydro-Input-Dateien gefunden. Weekly Constraints werden übersprungen.")
        return pd.DataFrame(columns=["week", "min_generated_mwh", "max_generated_mwh"]), {}

    weekly_agg: Dict[int, Dict[str, float]] = {}
    reference_caps: Dict[str, float] = {}

    for path in files:
        logger.info("Lade NO-Hydro-Constraints aus %s", path)
        tables = _read_hydro_input_tables(path)
        sheet_found = False

        for sheet_df in tables:
            sheet_text = []
            for row in sheet_df.iloc[:12, :8].itertuples(index=False, name=None):
                vals = [str(v) for v in row if pd.notna(v)]
                if vals:
                    sheet_text.extend(vals)
            sheet_text_join = " | ".join(sheet_text).lower()
            has_pump_storage_markers = (
                "pump" in sheet_text_join
                or "reservoir" in sheet_text_join
                or "cumulated" in sheet_text_join
            )
            if not has_pump_storage_markers:
                continue

            header_row, col_map = _find_label_row(sheet_df, ["Minimum Generated energy", "Maximum Generated energy"])
            if header_row is None or not col_map:
                continue

            week_col = None
            for probe_row_idx in [header_row, header_row + 1, header_row + 2]:
                if probe_row_idx >= len(sheet_df):
                    continue
                probe_row = sheet_df.iloc[probe_row_idx]
                for col_idx, value in enumerate(probe_row.tolist()):
                    if pd.isna(value):
                        continue
                    label = str(value).strip().lower()
                    if label in {"week", "day"}:
                        week_col = col_idx
                        break
                if week_col is not None:
                    break

            if week_col is None:
                continue

            min_col = col_map["Minimum Generated energy"]
            max_col = col_map["Maximum Generated energy"]
            weekly_values_found = False

            for row_idx in range(header_row + 1, len(sheet_df)):
                row = sheet_df.iloc[row_idx]
                week_val = _coerce_numeric(row.iloc[week_col])
                if week_val is None:
                    continue
                week = int(week_val)
                if week < 1:
                    continue

                min_mwh = _coerce_numeric(row.iloc[min_col])
                max_mwh = _coerce_numeric(row.iloc[max_col])
                if min_mwh is None and max_mwh is None:
                    continue

                if min_mwh is not None:
                    min_mwh *= 1000.0
                if max_mwh is not None:
                    max_mwh *= 1000.0

                weekly_agg.setdefault(week, {"min_generated_mwh": 0.0, "max_generated_mwh": 0.0})
                if min_mwh is not None:
                    weekly_agg[week]["min_generated_mwh"] += min_mwh
                if max_mwh is not None:
                    weekly_agg[week]["max_generated_mwh"] += max_mwh

                weekly_values_found = True

            if not weekly_values_found:
                continue

            turbine_cap = _find_numeric_value(sheet_df, ["Reference Total turbining capacity", "(T) Reference Total turbining capacity"])
            if turbine_cap is not None:
                reference_caps[path.stem] = float(turbine_cap)
                logger.info("  Referenz-Turbinenleistung %s: %.1f MW", path.name, turbine_cap)

            sheet_found = True
            break

        if not sheet_found:
            logger.info("Keine Weekly-Header in %s gefunden, überspringe Datei", path)

    if not weekly_agg:
        logger.info("Keine Weekly-Values gefunden. Weekly Constraints werden übersprungen.")
        return pd.DataFrame(columns=["week", "min_generated_mwh", "max_generated_mwh"]), reference_caps

    weekly_df = pd.DataFrame([
        {"week": week, "min_generated_mwh": values["min_generated_mwh"], "max_generated_mwh": values["max_generated_mwh"]}
        for week, values in sorted(weekly_agg.items())
    ])
    logger.info("Aggregierte NO-Hydro-Weekly-Constraints geladen für %d Wochen", len(weekly_df))
    return weekly_df, reference_caps


def load_no_hydro_reservoir_capacity(search_dir: Path, codes: Optional[List[str]] = None) -> Tuple[float, Dict[str, float]]:
    """Liest die aufsummierte Speicherenergie-Kapazität aus den norwegischen Dateien."""
    files = _find_no_hydro_files(search_dir, codes=codes)
    reservoir_caps: Dict[str, float] = {}
    total_gwh = 0.0

    for path in files:
        tables = _read_hydro_input_tables(path)
        reservoir_cap = None
        for sheet_df in tables:
            sheet_text = []
            for row in sheet_df.iloc[:12, :8].itertuples(index=False, name=None):
                vals = [str(v) for v in row if pd.notna(v)]
                if vals:
                    sheet_text.extend(vals)
            sheet_text_join = " | ".join(sheet_text).lower()
            if "pump" not in sheet_text_join and "reservoir" not in sheet_text_join:
                continue
            for row_idx, row in sheet_df.iterrows():
                row_vals = [str(val).strip().lower() if pd.notna(val) else "" for val in row.tolist()]
                for col_idx, val in enumerate(row_vals):
                    if "cumulated (upper or head) reservoir capacity" in val or "reservoir capacity" in val:
                        for candidate_idx in range(col_idx + 1, min(len(row), col_idx + 4)):
                            numeric = _coerce_numeric(sheet_df.iloc[row_idx, candidate_idx])
                            if numeric is not None:
                                reservoir_cap = numeric
                                break
                        if reservoir_cap is not None:
                            break
                if reservoir_cap is not None:
                    break
            if reservoir_cap is not None:
                break

        if reservoir_cap is not None:
            reservoir_caps[path.stem] = float(reservoir_cap)
            total_gwh += float(reservoir_cap)
            logger.info("  Speicherenergie %s: %.1f GWh", path.name, reservoir_cap)

    if not reservoir_caps:
        logger.info("Keine Speicherenergie-Kapazitäten gefunden; PSOpen_NO behält die vorhandene Kapazität bei.")
    return total_gwh, reservoir_caps


def map_weeks_to_snapshots(n: pypsa.Network) -> Dict[int, List[pd.Timestamp]]:
    """Mappt die Snapshot-Zeiten in 168h-Blöcke auf Wochen 1..N."""
    snapshots = list(n.snapshots)
    week_map: Dict[int, List[pd.Timestamp]] = {}
    week_idx = 1
    start = 0
    while start < len(snapshots):
        end = min(start + 168, len(snapshots))
        week_map[week_idx] = snapshots[start:end]
        start = end
        week_idx += 1

    logger.info("Snapshot-zu-Woche-Mapping erzeugt für %d Wochen", len(week_map))
    return week_map


def adjust_storage_capacity(n: pypsa.Network, reservoir_capacity_gwh: float, storage_name: str = "PSOpen_NO") -> None:
    """Passt die Speicherkapazität eines StorageUnits an die aggregierte Reservoir-Energie an."""
    if storage_name not in n.storage_units.index:
        logger.warning("%s nicht im Netz gefunden; Speicherkapazität wird nicht angepasst.", storage_name)
        return

    current_p_nom = float(n.storage_units.at[storage_name, "p_nom"])
    current_max_hours = float(n.storage_units.at[storage_name, "max_hours"])
    if current_p_nom <= 0:
        logger.warning("%s p_nom <= 0; Speicherkapazität wird nicht angepasst.", storage_name)
        return

    if reservoir_capacity_gwh <= 0:
        logger.warning("Aggregierte Speicherenergie ist <= 0; %s behält die bisherige Kapazität bei.", storage_name)
        return

    new_max_hours = (reservoir_capacity_gwh * 1000.0) / current_p_nom
    n.storage_units.at[storage_name, "max_hours"] = new_max_hours
    if "e_nom" in n.storage_units.columns:
        n.storage_units.at[storage_name, "e_nom"] = current_p_nom * new_max_hours

    logger.info(
        "%s Speicherenergie angepasst: vorher %.1f h, neu %.1f h (%.1f GWh, p_nom=%.1f MW)",
        storage_name,
        current_max_hours,
        new_max_hours,
        reservoir_capacity_gwh,
        current_p_nom,
    )


def add_weekly_flow_constraints(
    n: pypsa.Network,
    weekly_df: pd.DataFrame,
    week_snapshot_map: Dict[int, List[pd.Timestamp]],
    enforce_min_generation: bool = True,
    storage_name: str = "PSOpen_NO",
):
    """Ergänzt wöchentliche Max-/Min-Dispatch-Constraints für eine StorageUnit."""
    if storage_name not in n.storage_units.index:
        logger.warning("%s nicht im Netz gefunden; Weekly-Constraints werden übersprungen.", storage_name)
        return

    if "StorageUnit-p_dispatch" not in n.model.variables:
        logger.warning("StorageUnit-p_dispatch Variable nicht im Modell gefunden; Weekly-Constraints werden übersprungen.")
        return

    p_dispatch = n.model.variables["StorageUnit-p_dispatch"]
    applied = []

    for _, row in weekly_df.iterrows():
        week = int(row["week"])
        if week not in week_snapshot_map:
            continue

        snapshots = week_snapshot_map[week]
        if not snapshots:
            continue

        dispatch_expr = p_dispatch.sel(name=storage_name, snapshot=snapshots).sum("snapshot")
        min_mwh = float(row.get("min_generated_mwh", 0.0))
        max_mwh = float(row.get("max_generated_mwh", 0.0))

        if enforce_min_generation and min_mwh > 0:
            n.model.add_constraints(dispatch_expr >= min_mwh, name=f"NOHydroFix-min-{week}")
        if max_mwh > 0:
            n.model.add_constraints(dispatch_expr <= max_mwh, name=f"NOHydroFix-max-{week}")

        logger.info(
            "Weekly-Constraint Woche %d: min=%.1f MWh, max=%.1f MWh, snapshots=%d",
            week,
            min_mwh,
            max_mwh,
            len(snapshots),
        )
        applied.append(week)

    if applied:
        logger.info("%d Weekly-Constraints für %s gesetzt", len(applied), storage_name)


def write_weekly_check(
    n: pypsa.Network,
    weekly_df: pd.DataFrame,
    week_snapshot_map: Dict[int, List[pd.Timestamp]],
    out_dir: Path,
    storage_name: str = "PSOpen_NO",
    file_name: str = "NO_hydro_fix_weekly_check.csv",
) -> None:
    """Schreibt eine CSV mit tatsächlicher Wochen-Erzeugung und Grenzen."""
    if storage_name not in n.storage_units_t.p_dispatch.columns:
        return

    actual_dispatch = n.storage_units_t.p_dispatch[storage_name]
    rows = []
    for _, row in weekly_df.iterrows():
        week = int(row["week"])
        snapshots = week_snapshot_map.get(week, [])
        if not snapshots:
            continue
        actual_mwh = float(actual_dispatch.loc[snapshots].sum())
        rows.append({
            "week": week,
            "actual_generated_mwh": actual_mwh,
            "min_generated_mwh": float(row.get("min_generated_mwh", 0.0)),
            "max_generated_mwh": float(row.get("max_generated_mwh", 0.0)),
        })

    if rows:
        out_path = out_dir / file_name
        pd.DataFrame(rows).to_csv(out_path, index=False)
        logger.info("Weekly-Check gespeichert in %s", out_path)


def load_reservoir_level_constraints(
    search_dir: Path,
    codes: Optional[List[str]] = None,
    use_technical_bounds: bool = False,
) -> Tuple[pd.DataFrame, Dict[str, float]]:
    """
    Liest wöchentliche Reservoir-Level-Constraints aus PEMMDB-Dateien und
    aggregiert sie KAPAZITÄTSGEWICHTET über alle norwegischen Bietzonen.

    Gibt (level_df, zone_caps_gwh) zurück:
    - level_df: Spalten week, min_level_ratio, max_level_ratio, min_level_mwh, max_level_mwh
    - zone_caps_gwh: Reservoir-Kapazität pro Datei in GWh (für Diagnose)

    Hinweis: Es existiert in denselben Dateien ab Spalte ~377 auch ein
    klimajahr-spezifischer Block mit denselben Constraint-Typen. Dieser wird
    hier NICHT gelesen (Uniform Constraints in Cols 11/12 werden verwendet).
    """
    min_label = "Minimum Reservoir level, technical" if use_technical_bounds else "Minimum Reservoir level, historical"
    max_label = "Maximum Reservoir level, technical" if use_technical_bounds else "Maximum Reservoir level, historical"

    files = _find_no_hydro_files(search_dir, codes=codes)
    if not files:
        logger.warning("Keine NO-Hydro-Dateien für Reservoir-Level-Constraints gefunden.")
        return pd.DataFrame(columns=["week", "min_level_ratio", "max_level_ratio", "min_level_mwh", "max_level_mwh"]), {}

    # {week: {"weighted_min": float, "weighted_max": float, "total_cap": float}}
    agg: Dict[int, Dict[str, float]] = {}
    zone_caps_gwh: Dict[str, float] = {}

    for path in files:
        logger.info("Lade Reservoir-Level-Constraints aus %s", path)
        tables = _read_hydro_input_tables(path)

        for sheet_df in tables:
            sheet_text = " | ".join(
                str(v) for row in sheet_df.iloc[:12, :8].itertuples(index=False, name=None)
                for v in row if pd.notna(v)
            ).lower()
            if "pump" not in sheet_text and "reservoir" not in sheet_text:
                continue

            # Reservoir-Kapazität aus Zeile 6 (col 2)
            reservoir_cap_gwh: Optional[float] = None
            for row_idx, row in sheet_df.iterrows():
                row_vals = [str(v).strip().lower() if pd.notna(v) else "" for v in row.tolist()]
                for col_idx, val in enumerate(row_vals):
                    if "cumulated" in val and "reservoir" in val:
                        for ci in range(col_idx + 1, min(len(row), col_idx + 5)):
                            num = _coerce_numeric(sheet_df.iloc[row_idx, ci])
                            if num is not None and num > 0:
                                reservoir_cap_gwh = num
                                break
                    if reservoir_cap_gwh is not None:
                        break
                if reservoir_cap_gwh is not None:
                    break

            if reservoir_cap_gwh is None or reservoir_cap_gwh <= 0:
                logger.warning("Keine Reservoir-Kapazität in %s gefunden; überspringe.", path.name)
                continue

            zone_caps_gwh[path.stem] = reservoir_cap_gwh
            logger.info("  Reservoir-Kapazität %s: %.1f GWh", path.name, reservoir_cap_gwh)

            # Level-Constraint-Spalten finden
            header_row, col_map = _find_label_row(sheet_df, [min_label, max_label])
            if header_row is None or not col_map:
                logger.warning("  Level-Header '%s' nicht in %s gefunden.", min_label, path.name)
                continue

            week_col = None
            for probe_idx in [header_row, header_row + 1, header_row + 2]:
                if probe_idx >= len(sheet_df):
                    continue
                for ci, val in enumerate(sheet_df.iloc[probe_idx].tolist()):
                    if pd.notna(val) and str(val).strip().lower() in {"week", "day"}:
                        week_col = ci
                        break
                if week_col is not None:
                    break

            if week_col is None:
                logger.warning("  Keine 'Week'-Spalte in %s gefunden.", path.name)
                continue

            min_col = col_map[min_label]
            max_col = col_map[max_label]

            for row_idx in range(header_row + 1, len(sheet_df)):
                row = sheet_df.iloc[row_idx]
                week_val = _coerce_numeric(row.iloc[week_col])
                if week_val is None:
                    continue
                week = int(week_val)
                if week < 1:
                    continue

                min_ratio = _coerce_numeric(row.iloc[min_col])
                max_ratio = _coerce_numeric(row.iloc[max_col])
                if min_ratio is None and max_ratio is None:
                    continue

                entry = agg.setdefault(week, {"weighted_min": 0.0, "weighted_max": 0.0, "total_cap": 0.0})
                if min_ratio is not None:
                    entry["weighted_min"] += min_ratio * reservoir_cap_gwh
                if max_ratio is not None:
                    entry["weighted_max"] += max_ratio * reservoir_cap_gwh
                entry["total_cap"] += reservoir_cap_gwh

            break  # erstes passendes Sheet pro Datei

    if not agg:
        logger.warning("Keine Reservoir-Level-Daten geladen.")
        return pd.DataFrame(columns=["week", "min_level_ratio", "max_level_ratio", "min_level_mwh", "max_level_mwh"]), zone_caps_gwh

    total_cap_gwh = sum(zone_caps_gwh.values())
    total_cap_mwh = total_cap_gwh * 1000.0

    rows = []
    for week in sorted(agg):
        entry = agg[week]
        cap = entry["total_cap"]
        if cap <= 0:
            continue
        min_r = entry["weighted_min"] / cap
        max_r = entry["weighted_max"] / cap
        rows.append({
            "week": week,
            "min_level_ratio": min_r,
            "max_level_ratio": max_r,
            "min_level_mwh": min_r * total_cap_mwh,
            "max_level_mwh": max_r * total_cap_mwh,
        })
        logger.info(
            "  Reservoir-Level Woche %d: min=%.1f%% (%.0f MWh), max=%.1f%% (%.0f MWh)",
            week, min_r * 100, min_r * total_cap_mwh, max_r * 100, max_r * total_cap_mwh,
        )

    level_df = pd.DataFrame(rows)
    logger.info(
        "Reservoir-Level-Constraints geladen: %d Wochen, Gesamtkapazität=%.1f GWh (%s)",
        len(level_df), total_cap_gwh, "technical" if use_technical_bounds else "historical",
    )
    return level_df, zone_caps_gwh


def add_reservoir_level_constraints(
    n: pypsa.Network,
    level_df: pd.DataFrame,
    week_snapshot_map: Dict[int, List[pd.Timestamp]],
    storage_name: str = "PSOpen_NO",
) -> None:
    """
    Setzt wöchentliche State-of-Charge-Constraints für eine StorageUnit.

    Constraint-Logik: SoC am LETZTEN Snapshot jeder Woche muss im erlaubten
    Band [min_level_mwh, max_level_mwh] liegen. Die letzte Stunde der Woche
    entspricht dem 'Anfang der Folgewoche' in der PEMMDB-Terminologie.
    Eine Constraint pro Wochenende (nicht pro Stunde) hält das LP klein.
    """
    if storage_name not in n.storage_units.index:
        logger.warning("%s nicht im Netz; Reservoir-Level-Constraints übersprungen.", storage_name)
        return

    if "StorageUnit-state_of_charge" not in n.model.variables:
        logger.warning("StorageUnit-state_of_charge nicht im Modell; Constraints übersprungen.")
        return

    soc_var = n.model.variables["StorageUnit-state_of_charge"]

    # Startfüllstand prüfen (nur relevant wenn cyclic_state_of_charge=False)
    p_nom = float(n.storage_units.at[storage_name, "p_nom"])
    max_hours = float(n.storage_units.at[storage_name, "max_hours"])
    total_cap_mwh = p_nom * max_hours
    is_cyclic = bool(n.storage_units.at[storage_name, "cyclic_state_of_charge"]) if "cyclic_state_of_charge" in n.storage_units.columns else True

    if not level_df.empty:
        week1_min_mwh = float(level_df.iloc[0]["min_level_mwh"])
        week1_min_pct = float(level_df.iloc[0]["min_level_ratio"]) * 100
        if is_cyclic:
            logger.info(
                "cyclic_state_of_charge=True: Startfüllstand wird endogen bestimmt "
                "(Woche-1-End-Minimum: %.1f%% = %.0f MWh).",
                week1_min_pct, week1_min_mwh,
            )
        else:
            soc_initial = float(n.storage_units.at[storage_name, "state_of_charge_initial"]) if "state_of_charge_initial" in n.storage_units.columns else float("nan")
            if not pd.isna(soc_initial):
                soc_initial_pct = soc_initial / total_cap_mwh * 100
                if soc_initial < week1_min_mwh:
                    logger.warning(
                        "WARNUNG: Startfüllstand %.0f MWh (%.1f%%) liegt UNTER der "
                        "Woche-1-Mindestgrenze %.0f MWh (%.1f%%). "
                        "Prüfe state_of_charge_initial.",
                        soc_initial, soc_initial_pct, week1_min_mwh, week1_min_pct,
                    )
                else:
                    logger.info(
                        "Startfüllstand %.0f MWh (%.1f%%) liegt im erlaubten Band (Woche-1-Min: %.1f%%).",
                        soc_initial, soc_initial_pct, week1_min_pct,
                    )

    applied = []
    for _, row in level_df.iterrows():
        week = int(row["week"])
        snapshots = week_snapshot_map.get(week, [])
        if not snapshots:
            continue

        last_ts = snapshots[-1]
        soc_expr = soc_var.sel(name=storage_name, snapshot=last_ts)
        min_mwh = float(row["min_level_mwh"])
        max_mwh = float(row["max_level_mwh"])

        if min_mwh > 0:
            n.model.add_constraints(soc_expr >= min_mwh, name=f"NOReservoirLevel-min-{week}")
        if max_mwh > 0:
            n.model.add_constraints(soc_expr <= max_mwh, name=f"NOReservoirLevel-max-{week}")

        applied.append(week)

    if applied:
        logger.info("%d Reservoir-Level-Constraints für %s gesetzt", len(applied), storage_name)


def write_reservoir_level_check(
    n: pypsa.Network,
    level_df: pd.DataFrame,
    week_snapshot_map: Dict[int, List[pd.Timestamp]],
    out_dir: Path,
    storage_name: str = "PSOpen_NO",
    file_name: str = "NO_hydro_fix_reservoir_level_check.csv",
) -> None:
    """Schreibt CSV mit tatsächlichem Füllstand am Wochenende vs. erlaubtem Band."""
    if storage_name not in n.storage_units_t.state_of_charge.columns:
        logger.warning("state_of_charge für %s nicht in Ergebnissen.", storage_name)
        return

    p_nom = float(n.storage_units.at[storage_name, "p_nom"])
    max_hours = float(n.storage_units.at[storage_name, "max_hours"])
    total_cap_mwh = p_nom * max_hours

    soc_ts = n.storage_units_t.state_of_charge[storage_name]
    rows = []
    for _, row in level_df.iterrows():
        week = int(row["week"])
        snapshots = week_snapshot_map.get(week, [])
        if not snapshots:
            continue
        actual_mwh = float(soc_ts.loc[snapshots[-1]])
        actual_pct = actual_mwh / total_cap_mwh * 100 if total_cap_mwh > 0 else float("nan")
        rows.append({
            "week": week,
            "actual_soc_mwh": actual_mwh,
            "actual_soc_pct": actual_pct,
            "min_level_mwh": float(row["min_level_mwh"]),
            "min_level_pct": float(row["min_level_ratio"]) * 100,
            "max_level_mwh": float(row["max_level_mwh"]),
            "max_level_pct": float(row["max_level_ratio"]) * 100,
        })

    if rows:
        out_path = out_dir / file_name
        pd.DataFrame(rows).to_csv(out_path, index=False)
        logger.info("Reservoir-Level-Check gespeichert in %s", out_path)


def build_no_hydro_extra_functionality(
    n: pypsa.Network,
    level: str = "flow",
    search_dir: Path = config.ACCUMULATED_DIR,
) -> Tuple[callable, dict]:
    """
    Bereitet die NO-Hydro-Fix-Constraints für ein Netz vor und gibt eine
    extra_functionality-Callback-Funktion zurück (für n.optimize()), sowie
    die geladenen Wochendaten (für den späteren Check-CSV-Export nach dem Solve).

    level: "flow" (nur wöchentliche Erzeugungs-Constraints) oder
           "flow+level" (zusätzlich wöchentliche Reservoir-Level-Constraints).
    """
    zone_weekly_data: Dict[str, pd.DataFrame] = {}
    zone_level_data: Dict[str, pd.DataFrame] = {}
    zone_week_snapshot_map: Dict[str, Dict[int, List[pd.Timestamp]]] = {}

    for zone_config in HYDRO_ZONE_CONFIGS:
        zone_name = zone_config["zone"]
        storage_name = zone_config["storage_name"]
        codes = zone_config["codes"]

        weekly_df, reference_caps = load_no_hydro_weekly_constraints(search_dir, codes=codes)
        reservoir_capacity_gwh, _ = load_no_hydro_reservoir_capacity(search_dir, codes=codes)

        if reservoir_capacity_gwh > 0:
            adjust_storage_capacity(n, reservoir_capacity_gwh, storage_name=storage_name)
        else:
            logger.warning(
                "Keine aggregierte Reservoir-Kapazität verfügbar für %s; %s bleibt unverändert.",
                zone_name, storage_name,
            )

        week_snap_map = map_weeks_to_snapshots(n)
        zone_week_snapshot_map[zone_name] = week_snap_map

        if not weekly_df.empty:
            zone_weekly_data[zone_name] = weekly_df
            logger.info("Wöchentliche %s-Flow-Constraints vorbereitet: %d Wochen", zone_name, len(weekly_df))

        if level == "flow+level":
            level_df, _ = load_reservoir_level_constraints(
                search_dir, codes=codes,
                use_technical_bounds=(config.NO_HYDRO_RESERVOIR_BOUNDS == "technical"),
            )
            if not level_df.empty:
                zone_level_data[zone_name] = level_df
                logger.info("Reservoir-Level-Constraints vorbereitet für %s: %d Wochen", zone_name, len(level_df))

        if reference_caps:
            ref_total_mw = sum(reference_caps.values())
            model_p_nom = float(n.storage_units.at[storage_name, "p_nom"]) if storage_name in n.storage_units.index else float("nan")
            logger.warning(
                "%s-Hydro-Referenzleistung: Summe Bietzonen=%.1f MW, Modell %s p_nom=%.1f MW",
                zone_name, ref_total_mw, storage_name, model_p_nom,
            )

    def extra_functionality(n_network: pypsa.Network, snapshots):
        for zone_config in HYDRO_ZONE_CONFIGS:
            zone_name = zone_config["zone"]
            storage_name = zone_config["storage_name"]
            snap_map = zone_week_snapshot_map.get(zone_name, {})

            weekly_df = zone_weekly_data.get(zone_name)
            if weekly_df is not None and not weekly_df.empty:
                add_weekly_flow_constraints(
                    n_network, weekly_df, snap_map,
                    enforce_min_generation=True, storage_name=storage_name,
                )

            if level == "flow+level":
                level_df = zone_level_data.get(zone_name)
                if level_df is not None and not level_df.empty:
                    add_reservoir_level_constraints(n_network, level_df, snap_map, storage_name=storage_name)

    check_data = {
        "zone_weekly_data": zone_weekly_data,
        "zone_level_data": zone_level_data,
        "zone_week_snapshot_map": zone_week_snapshot_map,
    }
    return extra_functionality, check_data


def write_check_files(n: pypsa.Network, check_data: dict, out_dir: Path) -> None:
    """Schreibt die Weekly-/Reservoir-Level-Check-CSVs nach einem erfolgreichen Solve."""
    for zone_config in HYDRO_ZONE_CONFIGS:
        zone_name = zone_config["zone"]
        storage_name = zone_config["storage_name"]
        snap_map = check_data["zone_week_snapshot_map"].get(zone_name, {})

        weekly_df = check_data["zone_weekly_data"].get(zone_name)
        if weekly_df is not None and not weekly_df.empty:
            write_weekly_check(
                n, weekly_df, snap_map, out_dir,
                storage_name=storage_name,
                file_name=f"NO_hydro_fix_weekly_check_{zone_name}.csv",
            )

        level_df = check_data["zone_level_data"].get(zone_name)
        if level_df is not None and not level_df.empty:
            write_reservoir_level_check(
                n, level_df, snap_map, out_dir,
                storage_name=storage_name,
                file_name=f"NO_hydro_fix_reservoir_level_check_{zone_name}.csv",
            )
